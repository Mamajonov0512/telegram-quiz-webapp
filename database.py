import sqlite3
import json
import logging
import uuid
import random
from datetime import datetime
from typing import List, Dict, Any, Optional
from config import DB_PATH, ADMIN_IDS, DEFAULT_TEST_QUESTIONS_COUNT, DEFAULT_TEST_DURATION_MINUTES

logger = logging.getLogger(__name__)

def get_connection():
    conn = sqlite3.connect(str(DB_PATH), check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_connection()
    cursor = conn.cursor()
    
    # 1. Allowed Users / Whitelist table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS allowed_users (
        telegram_id INTEGER PRIMARY KEY,
        full_name TEXT DEFAULT '',
        username TEXT DEFAULT '',
        added_by INTEGER DEFAULT 0,
        added_at TEXT DEFAULT CURRENT_TIMESTAMP,
        is_active INTEGER DEFAULT 1,
        notes TEXT DEFAULT ''
    )
    """)

    # 2. Questions table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS questions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        question_text TEXT NOT NULL,
        option_a TEXT NOT NULL,
        option_b TEXT NOT NULL,
        option_c TEXT NOT NULL,
        option_d TEXT NOT NULL,
        correct_option TEXT NOT NULL,  -- 'A', 'B', 'C', or 'D'
        explanation TEXT DEFAULT '',
        category TEXT DEFAULT 'Umumiy',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )
    """)

    # 3. Test Results table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS test_results (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        telegram_id INTEGER NOT NULL,
        full_name TEXT DEFAULT '',
        username TEXT DEFAULT '',
        total_questions INTEGER NOT NULL,
        correct_answers INTEGER NOT NULL,
        wrong_answers INTEGER NOT NULL,
        score_percentage REAL NOT NULL,
        time_spent_seconds INTEGER NOT NULL,
        completed_at TEXT DEFAULT CURRENT_TIMESTAMP,
        answers_json TEXT DEFAULT '[]'
    )
    """)

    # 4. Settings table (key-value)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS settings (
        key TEXT PRIMARY KEY,
        value TEXT NOT NULL
    )
    """)

    # 5. Test Sessions table (tracks randomized questions and shuffled options per user)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS test_sessions (
        session_id TEXT PRIMARY KEY,
        telegram_id INTEGER NOT NULL,
        category TEXT DEFAULT 'Barchasi',
        questions_count INTEGER NOT NULL,
        duration_minutes INTEGER NOT NULL,
        session_data TEXT NOT NULL,
        started_at TEXT DEFAULT CURRENT_TIMESTAMP,
        is_submitted INTEGER DEFAULT 0
    )
    """)

    # 6. Sections table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS sections (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE NOT NULL,
        description TEXT DEFAULT '',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )
    """)

    # Migration for allowed_users: add allowed_sections if not exists
    cursor.execute("PRAGMA table_info(allowed_users)")
    cols = [r[1] for r in cursor.fetchall()]
    if "allowed_sections" not in cols:
        cursor.execute("ALTER TABLE allowed_users ADD COLUMN allowed_sections TEXT DEFAULT 'ALL'")

    # Seed sections table from existing distinct categories
    cursor.execute("""
    INSERT OR IGNORE INTO sections (name)
    SELECT DISTINCT category FROM questions 
    WHERE category IS NOT NULL AND TRIM(category) != ''
    """)

    # Set default settings if not exists
    defaults = {
        "whitelist_enabled": "true",  # Only allowed IDs can test
        "questions_per_test": str(DEFAULT_TEST_QUESTIONS_COUNT),
        "duration_minutes": str(DEFAULT_TEST_DURATION_MINUTES),
        "pass_percentage": "60",
        "shuffle_questions": "true",
        "shuffle_options": "true",
        "anti_cheat_enabled": "true",
        "category_filter_enabled": "true",
        "max_questions_limit": "500"
    }

    for k, v in defaults.items():
        cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES (?, ?)", (k, v))

    # Add admin IDs to whitelist automatically
    for admin_id in ADMIN_IDS:
        cursor.execute("""
        INSERT OR IGNORE INTO allowed_users (telegram_id, full_name, username, added_by, notes)
        VALUES (?, 'Administrator', 'admin', 0, 'Asosiy admin')
        """, (admin_id,))

    conn.commit()
    conn.close()
    logger.info("Database initialized successfully.")

# --- SETTINGS HELPERS ---
def get_setting(key: str, default: str = "") -> str:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM settings WHERE key = ?", (key,))
    row = cursor.fetchone()
    conn.close()
    return row["value"] if row else default

def set_setting(key: str, value: str):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (key, str(value)))
    conn.commit()
    conn.close()

def get_all_settings() -> Dict[str, str]:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT key, value FROM settings")
    rows = cursor.fetchall()
    conn.close()
    return {row["key"]: row["value"] for row in rows}

# --- ACCESS CONTROL / WHITELIST HELPERS ---
def is_user_allowed(telegram_id: int) -> bool:
    # If admin, always allowed
    if telegram_id in ADMIN_IDS:
        return True
    
    # If whitelist mode is disabled, anyone is allowed
    whitelist_mode = get_setting("whitelist_enabled", "true").lower() == "true"
    if not whitelist_mode:
        return True

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT is_active FROM allowed_users WHERE telegram_id = ?", (telegram_id,))
    row = cursor.fetchone()
    conn.close()
    return bool(row and row["is_active"] == 1)

def add_allowed_user(telegram_id: int, full_name: str = "", username: str = "", added_by: int = 0, notes: str = "", allowed_sections: Any = "ALL") -> bool:
    conn = get_connection()
    cursor = conn.cursor()
    try:
        val = "ALL"
        if isinstance(allowed_sections, list):
            if "ALL" in allowed_sections or "all" in allowed_sections:
                val = "ALL"
            else:
                val = json.dumps([str(s).strip() for s in allowed_sections if str(s).strip()], ensure_ascii=False)
        elif isinstance(allowed_sections, str):
            val = allowed_sections.strip() or "ALL"

        cursor.execute("""
        INSERT INTO allowed_users (telegram_id, full_name, username, added_by, added_at, is_active, notes, allowed_sections)
        VALUES (?, ?, ?, ?, ?, 1, ?, ?)
        ON CONFLICT(telegram_id) DO UPDATE SET
            full_name = excluded.full_name,
            username = excluded.username,
            is_active = 1,
            notes = excluded.notes,
            allowed_sections = excluded.allowed_sections
        """, (telegram_id, full_name, username, added_by, datetime.now().isoformat(), notes, val))
        conn.commit()
        return True
    except Exception as e:
        logger.error(f"Error adding allowed user: {e}")
        return False
    finally:
        conn.close()

def remove_allowed_user(telegram_id: int) -> bool:
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM allowed_users WHERE telegram_id = ?", (telegram_id,))
        conn.commit()
        return True
    except Exception as e:
        logger.error(f"Error removing user: {e}")
        return False
    finally:
        conn.close()

def toggle_user_active(telegram_id: int, is_active: bool) -> bool:
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("UPDATE allowed_users SET is_active = ? WHERE telegram_id = ?", (1 if is_active else 0, telegram_id))
        conn.commit()
        return True
    except Exception as e:
        logger.error(f"Error toggling user: {e}")
        return False
    finally:
        conn.close()

def get_allowed_users() -> List[Dict[str, Any]]:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM allowed_users ORDER BY added_at DESC")
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]

def get_user_allowed_sections(telegram_id: int) -> List[str]:
    """
    Returns list of section names this user is allowed to access.
    Admins always have access to ALL sections.
    """
    all_secs = [s["name"] for s in get_all_sections()]
    if telegram_id in ADMIN_IDS:
        return all_secs
    
    whitelist_mode = get_setting("whitelist_enabled", "true").lower() == "true"
    if not whitelist_mode:
        return all_secs

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT allowed_sections FROM allowed_users WHERE telegram_id = ?", (telegram_id,))
    row = cursor.fetchone()
    conn.close()

    if not row:
        return []
    
    raw = row["allowed_sections"]
    if not raw or str(raw).strip().upper() == "ALL":
        return all_secs
    
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            return [str(x).strip() for x in parsed if str(x).strip()]
    except Exception:
        pass
    
    return [p.strip() for p in str(raw).split(",") if p.strip()]

def set_user_allowed_sections(telegram_id: int, sections: Any) -> bool:
    """
    Sets allowed sections for a user.
    Can be 'ALL', or a list of section names, or a comma-separated string.
    """
    conn = get_connection()
    cursor = conn.cursor()
    try:
        val = "ALL"
        if isinstance(sections, list):
            if "ALL" in sections or "all" in sections:
                val = "ALL"
            else:
                val = json.dumps([str(s).strip() for s in sections if str(s).strip()], ensure_ascii=False)
        elif isinstance(sections, str):
            if sections.strip().upper() == "ALL":
                val = "ALL"
            else:
                parts = [p.strip() for p in sections.split(",") if p.strip()]
                val = json.dumps(parts, ensure_ascii=False)
        
        cursor.execute("UPDATE allowed_users SET allowed_sections = ? WHERE telegram_id = ?", (val, telegram_id))
        conn.commit()
        return True
    except Exception as e:
        logger.error(f"Error setting user sections: {e}")
        return False
    finally:
        conn.close()

# --- SECTIONS CRUD HELPERS ---
def get_all_sections() -> List[Dict[str, Any]]:
    conn = get_connection()
    cursor = conn.cursor()
    # Ensure any distinct category from questions is tracked in sections
    cursor.execute("""
    INSERT OR IGNORE INTO sections (name)
    SELECT DISTINCT category FROM questions 
    WHERE category IS NOT NULL AND TRIM(category) != ''
    """)
    conn.commit()

    cursor.execute("""
    SELECT s.id, s.name, s.description, s.created_at, COUNT(q.id) as count
    FROM sections s
    LEFT JOIN questions q ON q.category = s.name
    GROUP BY s.name
    ORDER BY count DESC, s.name ASC
    """)
    rows = cursor.fetchall()
    conn.close()
    return [{"id": r["id"], "name": r["name"], "category": r["name"], "description": r["description"], "count": r["count"], "created_at": r["created_at"]} for r in rows]

def add_section(name: str, description: str = "") -> bool:
    name = name.strip()
    if not name:
        return False
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO sections (name, description) VALUES (?, ?)", (name, description.strip()))
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    except Exception as e:
        logger.error(f"Error adding section: {e}")
        return False
    finally:
        conn.close()

def rename_section(old_name: str, new_name: str) -> bool:
    old_name = old_name.strip()
    new_name = new_name.strip()
    if not old_name or not new_name or old_name == new_name:
        return False
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("UPDATE sections SET name = ? WHERE name = ?", (new_name, old_name))
        cursor.execute("UPDATE questions SET category = ? WHERE category = ?", (new_name, old_name))
        # Update user permissions if specific section name stored
        cursor.execute("SELECT telegram_id, allowed_sections FROM allowed_users WHERE allowed_sections != 'ALL'")
        user_rows = cursor.fetchall()
        for u in user_rows:
            raw = u["allowed_sections"]
            if raw and old_name in raw:
                try:
                    parsed = json.loads(raw)
                    if isinstance(parsed, list):
                        updated = [new_name if x == old_name else x for x in parsed]
                        cursor.execute("UPDATE allowed_users SET allowed_sections = ? WHERE telegram_id = ?", 
                                       (json.dumps(updated, ensure_ascii=False), u["telegram_id"]))
                except Exception:
                    pass
        conn.commit()
        return True
    except Exception as e:
        logger.error(f"Error renaming section: {e}")
        return False
    finally:
        conn.close()

def delete_section(name: str, fallback_section: str = "Umumiy") -> bool:
    name = name.strip()
    if not name:
        return False
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM sections WHERE name = ?", (name,))
        if name != fallback_section:
            cursor.execute("INSERT OR IGNORE INTO sections (name) VALUES (?)", (fallback_section,))
            cursor.execute("UPDATE questions SET category = ? WHERE category = ?", (fallback_section, name))
        conn.commit()
        return True
    except Exception as e:
        logger.error(f"Error deleting section: {e}")
        return False
    finally:
        conn.close()

# --- QUESTIONS HELPERS ---
def add_question(question_text: str, option_a: str, option_b: str, option_c: str, option_d: str, 
                 correct_option: str, explanation: str = "", category: str = "Umumiy") -> int:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT INTO questions (question_text, option_a, option_b, option_c, option_d, correct_option, explanation, category)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (question_text.strip(), option_a.strip(), option_b.strip(), option_c.strip(), option_d.strip(), 
          correct_option.strip().upper(), explanation.strip(), category.strip()))
    conn.commit()
    q_id = cursor.lastrowid
    conn.close()
    return q_id

def bulk_add_questions(questions_list: List[Dict[str, Any]]) -> int:
    conn = get_connection()
    cursor = conn.cursor()
    count = 0
    for q in questions_list:
        try:
            cursor.execute("""
            INSERT INTO questions (question_text, option_a, option_b, option_c, option_d, correct_option, explanation, category)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                q.get("question_text", "").strip(),
                q.get("option_a", "").strip(),
                q.get("option_b", "").strip(),
                q.get("option_c", "").strip(),
                q.get("option_d", "").strip(),
                q.get("correct_option", "A").strip().upper(),
                q.get("explanation", "").strip(),
                q.get("category", "Umumiy").strip()
            ))
            count += 1
        except Exception as e:
            logger.error(f"Error inserting question: {e}")
    conn.commit()
    conn.close()
    return count

def get_all_questions(limit: int = 500, offset: int = 0) -> List[Dict[str, Any]]:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM questions ORDER BY id DESC LIMIT ? OFFSET ?", (limit, offset))
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]

def get_questions_count() -> int:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) as cnt FROM questions")
    row = cursor.fetchone()
    conn.close()
    return row["cnt"] if row else 0

def clear_all_questions() -> bool:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM questions")
    conn.commit()
    conn.close()
    return True

def delete_question(question_id: int) -> bool:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM questions WHERE id = ?", (question_id,))
    conn.commit()
    conn.close()
    return True

def update_question_category(question_id: int, new_category: str) -> bool:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE questions SET category = ? WHERE id = ?", (new_category.strip(), question_id))
    conn.commit()
    conn.close()
    return True

def update_all_questions_category(new_category: str, old_category: Optional[str] = None) -> int:
    conn = get_connection()
    cursor = conn.cursor()
    if old_category:
        cursor.execute("UPDATE questions SET category = ? WHERE category = ?", (new_category.strip(), old_category.strip()))
    else:
        cursor.execute("UPDATE questions SET category = ?", (new_category.strip(),))
    updated_rows = cursor.rowcount
    conn.commit()
    conn.close()
    return updated_rows

def get_categories(telegram_id: Optional[int] = None) -> List[Dict[str, Any]]:
    """
    Returns list of all sections with question count for each.
    Always returns both 'category' and 'name' fields to ensure frontends never see undefined.
    If telegram_id is provided, filters strictly to sections permitted for that user.
    """
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT OR IGNORE INTO sections (name)
    SELECT DISTINCT category FROM questions 
    WHERE category IS NOT NULL AND TRIM(category) != ''
    """)
    conn.commit()

    cursor.execute("""
    SELECT s.name, COUNT(q.id) as count 
    FROM sections s
    LEFT JOIN questions q ON q.category = s.name
    GROUP BY s.name 
    ORDER BY count DESC, s.name ASC
    """)
    rows = cursor.fetchall()
    conn.close()

    all_sections = [
        {"category": r["name"], "name": r["name"], "count": r["count"]} 
        for r in rows
    ]

    if telegram_id:
        allowed = get_user_allowed_sections(telegram_id)
        return [s for s in all_sections if s["name"] in allowed]

    return all_sections

def get_test_questions(count: int = 50, shuffle: bool = True, category: Optional[str] = None, telegram_id: Optional[int] = None) -> List[Dict[str, Any]]:
    conn = get_connection()
    cursor = conn.cursor()
    count = min(max(1, count), 500)
    order_clause = "ORDER BY RANDOM()" if shuffle else "ORDER BY id ASC"
    
    allowed = get_user_allowed_sections(telegram_id) if telegram_id else None

    if category and category.lower() not in ["barchasi", "all", "umumiy"]:
        if allowed is not None and telegram_id not in ADMIN_IDS and category not in allowed:
            conn.close()
            return []
        cursor.execute(f"SELECT * FROM questions WHERE category = ? {order_clause} LIMIT ?", (category, count))
    else:
        if allowed is not None and telegram_id not in ADMIN_IDS:
            if not allowed:
                conn.close()
                return []
            placeholders = ",".join("?" * len(allowed))
            cursor.execute(f"SELECT * FROM questions WHERE category IN ({placeholders}) {order_clause} LIMIT ?", (*allowed, count))
        else:
            cursor.execute(f"SELECT * FROM questions {order_clause} LIMIT ?", (count,))
            
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]

def create_test_session(telegram_id: int, category: Optional[str] = None, count: Optional[int] = None) -> Dict[str, Any]:
    """
    Initializes a test session with:
    - Selected section filter (controlled by category_filter_enabled setting)
    - User section permission check (users only take tests from permitted sections)
    - Specified question count limit (capped by max_questions_limit, max 500)
    - Randomized question order (controlled by shuffle_questions setting)
    - Randomized options order (controlled by shuffle_options setting)
    """
    max_limit = int(get_setting("max_questions_limit", "500"))
    max_limit = min(max(1, max_limit), 500)

    default_count = int(get_setting("questions_per_test", str(DEFAULT_TEST_QUESTIONS_COUNT)))
    target_count = count if (count and count > 0) else default_count
    target_count = min(max(1, target_count), max_limit)

    base_duration = int(get_setting("duration_minutes", str(DEFAULT_TEST_DURATION_MINUTES)))
    duration_minutes = target_count if (count and count > 0) else base_duration
    
    shuffle_q_enabled = get_setting("shuffle_questions", "true").lower() == "true"
    order_clause = "ORDER BY RANDOM()" if shuffle_q_enabled else "ORDER BY id ASC"

    # Enforce section permissions
    allowed_sections = get_user_allowed_sections(telegram_id)
    if not allowed_sections and telegram_id not in ADMIN_IDS:
        return {"error": "Sizga test topshirish uchun birorta ham bo'limga ruxsat berilmagan. Administratorga murojaat qiling.", "questions": [], "count": 0, "session_id": ""}

    conn = get_connection()
    cursor = conn.cursor()

    cat_filter_enabled = get_setting("category_filter_enabled", "true").lower() == "true"
    if cat_filter_enabled and category and category.lower() not in ["barchasi", "all"]:
        if telegram_id not in ADMIN_IDS and category not in allowed_sections:
            conn.close()
            return {"error": f"Sizga '{category}' bo'limidagi testlarni ishlashga ruxsat berilmagan!", "questions": [], "count": 0, "session_id": ""}
        cursor.execute(f"SELECT * FROM questions WHERE category = ? {order_clause} LIMIT ?", (category, target_count))
    else:
        if telegram_id not in ADMIN_IDS:
            placeholders = ",".join("?" * len(allowed_sections))
            cursor.execute(f"SELECT * FROM questions WHERE category IN ({placeholders}) {order_clause} LIMIT ?", (*allowed_sections, target_count))
        else:
            cursor.execute(f"SELECT * FROM questions {order_clause} LIMIT ?", (target_count,))

    raw_questions = [dict(r) for r in cursor.fetchall()]
    conn.close()

    if not raw_questions:
        return {"error": "Tanlangan bo'limda savollar mavjud emas", "questions": [], "count": 0, "session_id": ""}

    shuffle_options_enabled = get_setting("shuffle_options", "true").lower() == "true"
    
    session_id = str(uuid.uuid4())
    client_questions = []
    session_questions_map = {}

    for idx, q in enumerate(raw_questions, start=1):
        original_correct = q["correct_option"].upper().strip()
        
        # Raw options list
        raw_options = [
            ("A", q["option_a"]),
            ("B", q["option_b"]),
            ("C", q["option_c"]),
            ("D", q["option_d"]),
        ]

        if shuffle_options_enabled:
            # Shuffle options order randomly
            random.shuffle(raw_options)

        # Assign new A, B, C, D keys and track where original correct answer landed
        shuffled_options = {}
        correct_shown_letter = "A"

        for new_letter, (orig_letter, opt_text) in zip(["A", "B", "C", "D"], raw_options):
            shuffled_options[new_letter] = opt_text
            if orig_letter == original_correct:
                correct_shown_letter = new_letter

        # Save to session map for grading
        session_questions_map[str(q["id"])] = {
            "id": q["id"],
            "question_text": q["question_text"],
            "category": q.get("category", "Umumiy"),
            "options": shuffled_options,
            "correct_shown_letter": correct_shown_letter,
            "explanation": q.get("explanation", "")
        }

        # Client-facing question (SAFE: never contains correct answer or explanation)
        client_questions.append({
            "id": q["id"],
            "index": idx,
            "question": q["question_text"],
            "options": shuffled_options,
            "category": q.get("category", "Umumiy")
        })

    # Save test session
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT INTO test_sessions 
    (session_id, telegram_id, category, questions_count, duration_minutes, session_data, started_at, is_submitted)
    VALUES (?, ?, ?, ?, ?, ?, ?, 0)
    """, (
        session_id,
        telegram_id,
        category or "Barchasi",
        len(client_questions),
        duration_minutes,
        json.dumps(session_questions_map, ensure_ascii=False),
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ))
    conn.commit()
    conn.close()

    return {
        "session_id": session_id,
        "questions": client_questions,
        "count": len(client_questions),
        "duration_minutes": duration_minutes,
        "category": category or "Barchasi"
    }

def evaluate_session_submission(session_id: str, submitted_answers: Dict[str, str], time_spent_seconds: int,
                                full_name: str = "", username: str = "") -> Optional[Dict[str, Any]]:
    """
    Accurately evaluates submitted answers against the shuffled session options.
    """
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM test_sessions WHERE session_id = ?", (session_id,))
    session_row = cursor.fetchone()
    
    if not session_row:
        conn.close()
        return None

    telegram_id = session_row["telegram_id"]
    session_data = json.loads(session_row["session_data"])

    correct_count = 0
    wrong_count = 0
    detailed_review = []
    total_questions = len(session_data)

    for qid_str, q_info in session_data.items():
        user_ans = (submitted_answers.get(qid_str) or "").upper().strip()
        correct_ans = q_info["correct_shown_letter"].upper().strip()
        is_correct = (user_ans == correct_ans and bool(user_ans))

        if is_correct:
            correct_count += 1
        else:
            wrong_count += 1

        detailed_review.append({
            "id": q_info["id"],
            "question": q_info["question_text"],
            "options": q_info["options"],
            "user_answer": user_ans,
            "correct_answer": correct_ans,
            "is_correct": is_correct,
            "explanation": q_info.get("explanation", ""),
            "category": q_info.get("category", "")
        })

    score_percentage = round((correct_count / total_questions) * 100, 1) if total_questions > 0 else 0
    pass_percentage = int(get_setting("pass_percentage", "60"))
    passed = score_percentage >= pass_percentage

    # Save to test_results
    review_json = json.dumps(detailed_review, ensure_ascii=False)
    cursor.execute("""
    INSERT INTO test_results 
    (telegram_id, full_name, username, total_questions, correct_answers, wrong_answers, score_percentage, time_spent_seconds, completed_at, answers_json)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (telegram_id, full_name, username, total_questions, correct_count, wrong_count,
          score_percentage, time_spent_seconds, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), review_json))
    
    # Mark session as submitted
    cursor.execute("UPDATE test_sessions SET is_submitted = 1 WHERE session_id = ?", (session_id,))
    conn.commit()
    conn.close()

    return {
        "score": correct_count,
        "total": total_questions,
        "wrong": wrong_count,
        "percentage": score_percentage,
        "passed": passed,
        "pass_percentage": pass_percentage,
        "time_spent_seconds": time_spent_seconds,
        "review": detailed_review
    }

# --- TEST RESULTS HELPERS ---
def save_test_result(telegram_id: int, full_name: str, username: str, total_questions: int,
                     correct_answers: int, wrong_answers: int, score_percentage: float,
                     time_spent_seconds: int, answers_json: str) -> int:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT INTO test_results 
    (telegram_id, full_name, username, total_questions, correct_answers, wrong_answers, score_percentage, time_spent_seconds, completed_at, answers_json)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (telegram_id, full_name, username, total_questions, correct_answers, wrong_answers,
          score_percentage, time_spent_seconds, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), answers_json))
    conn.commit()
    res_id = cursor.lastrowid
    conn.close()
    return res_id

def get_user_results(telegram_id: int, limit: int = 10) -> List[Dict[str, Any]]:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM test_results WHERE telegram_id = ? ORDER BY id DESC LIMIT ?", (telegram_id, limit))
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]

def get_all_results(limit: int = 100) -> List[Dict[str, Any]]:
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM test_results ORDER BY id DESC LIMIT ?", (limit,))
    rows = cursor.fetchall()
    conn.close()
    return [dict(row) for row in rows]

def get_stats() -> Dict[str, Any]:
    conn = get_connection()
    cursor = conn.cursor()
    
    cursor.execute("SELECT COUNT(*) as total_users FROM allowed_users")
    total_users = cursor.fetchone()["total_users"]

    cursor.execute("SELECT COUNT(*) as total_questions FROM questions")
    total_questions = cursor.fetchone()["total_questions"]

    cursor.execute("SELECT COUNT(*) as total_tests_taken FROM test_results")
    total_tests_taken = cursor.fetchone()["total_tests_taken"]

    cursor.execute("SELECT AVG(score_percentage) as avg_score FROM test_results")
    avg_row = cursor.fetchone()
    avg_score = round(avg_row["avg_score"] or 0, 1)

    conn.close()
    return {
        "total_users": total_users,
        "total_questions": total_questions,
        "total_tests_taken": total_tests_taken,
        "avg_score": avg_score
    }

# --- ADVANCED ANALYTICS FUNCTIONS ---

def get_detailed_analytics() -> Dict[str, Any]:
    """
    Computes comprehensive analytics across all test attempts:
    - Overall summary (participants, tests, avg score, pass rate, avg time)
    - Score distribution (0-39%, 40-59%, 60-79%, 80-100%)
    - Top participants leaderboard
    - Hardest questions (highest error rates)
    - Category accuracy analysis
    - Daily tests trend
    """
    conn = get_connection()
    cursor = conn.cursor()

    pass_percentage = int(get_setting("pass_percentage", "60"))

    # 1. Summary stats
    cursor.execute("""
    SELECT 
        COUNT(*) as total_tests,
        COUNT(DISTINCT telegram_id) as unique_users,
        AVG(score_percentage) as avg_score,
        MAX(score_percentage) as max_score,
        MIN(score_percentage) as min_score,
        AVG(time_spent_seconds) as avg_time_sec
    FROM test_results
    """)
    sum_row = cursor.fetchone()

    total_tests = sum_row["total_tests"] or 0
    unique_users = sum_row["unique_users"] or 0
    avg_score = round(sum_row["avg_score"] or 0, 1)
    max_score = round(sum_row["max_score"] or 0, 1)
    min_score = round(sum_row["min_score"] or 0, 1)
    avg_time_sec = int(sum_row["avg_time_sec"] or 0)

    if total_tests == 0:
        conn.close()
        return {
            "total_tests": 0,
            "unique_users": 0,
            "avg_score": 0,
            "max_score": 0,
            "min_score": 0,
            "avg_time_min": 0,
            "pass_count": 0,
            "fail_count": 0,
            "pass_rate": 0,
            "score_distribution": {"0-39": 0, "40-59": 0, "60-79": 0, "80-100": 0},
            "leaderboard": [],
            "hardest_questions": [],
            "category_performance": [],
            "daily_trend": []
        }

    # Pass vs Fail
    cursor.execute("SELECT COUNT(*) as passed FROM test_results WHERE score_percentage >= ?", (pass_percentage,))
    pass_count = cursor.fetchone()["passed"]
    fail_count = total_tests - pass_count
    pass_rate = round((pass_count / total_tests) * 100, 1)

    # 2. Score distribution
    cursor.execute("""
    SELECT 
        SUM(CASE WHEN score_percentage < 40 THEN 1 ELSE 0 END) as b1,
        SUM(CASE WHEN score_percentage >= 40 AND score_percentage < 60 THEN 1 ELSE 0 END) as b2,
        SUM(CASE WHEN score_percentage >= 60 AND score_percentage < 80 THEN 1 ELSE 0 END) as b3,
        SUM(CASE WHEN score_percentage >= 80 THEN 1 ELSE 0 END) as b4
    FROM test_results
    """)
    dist_row = cursor.fetchone()
    score_distribution = {
        "0-39": dist_row["b1"] or 0,
        "40-59": dist_row["b2"] or 0,
        "60-79": dist_row["b3"] or 0,
        "80-100": dist_row["b4"] or 0
    }

    # 3. Leaderboard (Top 10 users)
    cursor.execute("""
    SELECT 
        telegram_id,
        MAX(full_name) as full_name,
        MAX(username) as username,
        COUNT(*) as attempts_count,
        MAX(score_percentage) as best_score,
        ROUND(AVG(score_percentage), 1) as avg_score,
        ROUND(AVG(time_spent_seconds), 0) as avg_time
    FROM test_results
    GROUP BY telegram_id
    ORDER BY best_score DESC, avg_score DESC
    LIMIT 10
    """)
    leaderboard = [dict(row) for row in cursor.fetchall()]

    # 4. Daily trend (last 7 days)
    cursor.execute("""
    SELECT 
        SUBSTR(completed_at, 1, 10) as test_date,
        COUNT(*) as count,
        ROUND(AVG(score_percentage), 1) as day_avg
    FROM test_results
    GROUP BY test_date
    ORDER BY test_date DESC
    LIMIT 7
    """)
    daily_trend = [dict(row) for row in cursor.fetchall()]
    daily_trend.reverse()

    # 5. In-depth Question & Category Analysis from answers_json
    cursor.execute("SELECT answers_json FROM test_results")
    all_answers_rows = cursor.fetchall()
    conn.close()

    question_stats = {}  # q_id -> { "question": str, "category": str, "total": int, "wrong": int }
    category_stats = {}  # category -> { "total": int, "correct": int }

    for r in all_answers_rows:
        raw_json = r["answers_json"]
        if not raw_json:
            continue
        try:
            items = json.loads(raw_json)
            for item in items:
                qid = item.get("id")
                q_text = item.get("question", "")
                cat = item.get("category") or "Umumiy"
                is_correct = bool(item.get("is_correct"))

                # Question stats
                if qid not in question_stats:
                    question_stats[qid] = {
                        "id": qid,
                        "question": q_text,
                        "category": cat,
                        "total": 0,
                        "wrong": 0
                    }
                question_stats[qid]["total"] += 1
                if not is_correct:
                    question_stats[qid]["wrong"] += 1

                # Category stats
                if cat not in category_stats:
                    category_stats[cat] = {"category": cat, "total": 0, "correct": 0}
                category_stats[cat]["total"] += 1
                if is_correct:
                    category_stats[cat]["correct"] += 1
        except Exception:
            continue

    # Format hardest questions
    hardest_questions = []
    for q_data in question_stats.values():
        if q_data["total"] > 0:
            err_rate = round((q_data["wrong"] / q_data["total"]) * 100, 1)
            hardest_questions.append({
                "id": q_data["id"],
                "question": q_data["question"],
                "category": q_data["category"],
                "total_answered": q_data["total"],
                "wrong_count": q_data["wrong"],
                "error_rate": err_rate
            })

    # Sort hardest by error rate descending
    hardest_questions.sort(key=lambda x: (x["error_rate"], x["wrong_count"]), reverse=True)
    hardest_questions = hardest_questions[:10]

    # Format category performance
    category_performance = []
    for cat_data in category_stats.values():
        tot = cat_data["total"]
        cor = cat_data["correct"]
        acc = round((cor / tot) * 100, 1) if tot > 0 else 0
        category_performance.append({
            "category": cat_data["category"],
            "total_questions": tot,
            "correct_answers": cor,
            "accuracy_percentage": acc
        })
    category_performance.sort(key=lambda x: x["accuracy_percentage"])

    return {
        "total_tests": total_tests,
        "unique_users": unique_users,
        "avg_score": avg_score,
        "max_score": max_score,
        "min_score": min_score,
        "avg_time_min": round(avg_time_sec / 60, 1),
        "pass_count": pass_count,
        "fail_count": fail_count,
        "pass_rate": pass_rate,
        "score_distribution": score_distribution,
        "leaderboard": leaderboard,
        "hardest_questions": hardest_questions,
        "category_performance": category_performance,
        "daily_trend": daily_trend
    }

def get_user_analytics(telegram_id: int) -> Dict[str, Any]:
    """
    Computes personalized analytics for a single participant:
    - Attempt count, best score, average score
    - Category accuracy (strong & weak areas)
    - Trend of all tests taken
    """
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("""
    SELECT 
        COUNT(*) as total_attempts,
        MAX(score_percentage) as best_score,
        MIN(score_percentage) as min_score,
        ROUND(AVG(score_percentage), 1) as avg_score,
        ROUND(AVG(time_spent_seconds), 0) as avg_time_sec
    FROM test_results
    WHERE telegram_id = ?
    """, (telegram_id,))
    row = cursor.fetchone()

    total_attempts = row["total_attempts"] or 0
    if total_attempts == 0:
        conn.close()
        return {
            "total_attempts": 0,
            "best_score": 0,
            "avg_score": 0,
            "min_score": 0,
            "avg_time_min": 0,
            "category_performance": [],
            "history": []
        }

    # History
    cursor.execute("""
    SELECT id, score_percentage, correct_answers, total_questions, time_spent_seconds, completed_at, answers_json
    FROM test_results
    WHERE telegram_id = ?
    ORDER BY id ASC
    """, (telegram_id,))
    history_rows = cursor.fetchall()
    conn.close()

    # Category performance for this user
    category_stats = {}
    history_list = []

    for r in history_rows:
        history_list.append({
            "id": r["id"],
            "score": r["score_percentage"],
            "correct": r["correct_answers"],
            "total": r["total_questions"],
            "time_sec": r["time_spent_seconds"],
            "date": r["completed_at"]
        })

        raw_json = r["answers_json"]
        if raw_json:
            try:
                items = json.loads(raw_json)
                for item in items:
                    cat = item.get("category") or "Umumiy"
                    is_correct = bool(item.get("is_correct"))
                    if cat not in category_stats:
                        category_stats[cat] = {"category": cat, "total": 0, "correct": 0}
                    category_stats[cat]["total"] += 1
                    if is_correct:
                        category_stats[cat]["correct"] += 1
            except Exception:
                pass

    cat_list = []
    for c in category_stats.values():
        tot = c["total"]
        cor = c["correct"]
        acc = round((cor / tot) * 100, 1) if tot > 0 else 0
        status = "A'lo" if acc >= 80 else ("Yaxshi" if acc >= 60 else "Qayta takrorlang")
        cat_list.append({
            "category": c["category"],
            "total": tot,
            "correct": cor,
            "accuracy": acc,
            "status": status
        })
    cat_list.sort(key=lambda x: x["accuracy"])

    return {
        "total_attempts": total_attempts,
        "best_score": row["best_score"] or 0,
        "avg_score": row["avg_score"] or 0,
        "min_score": row["min_score"] or 0,
        "avg_time_min": round((row["avg_time_sec"] or 0) / 60, 1),
        "category_performance": cat_list,
        "history": history_list
    }

def generate_results_excel() -> bytes:
    """
    Generates a beautifully formatted Excel report (.xlsx) containing:
    1. Barcha Natijalar (All test results)
    2. Foydalanuvchilar Reytingi (Leaderboard)
    3. Qiyin Savollar Tahlili (Hardest questions)
    """
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    
    # Header styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # --- SHEET 1: TEST NATIJALARI ---
    ws1 = wb.active
    ws1.title = "Barcha Natijalar"

    headers1 = ["№", "Foydalanuvchi", "Telegram ID", "To'g'ri", "Noto'g'ri", "Jami", "Ball (%)", "Holat", "Vaqt (daq)", "Sana"]
    ws1.append(headers1)

    for col_idx, _ in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    results = get_all_results(limit=1000)
    for idx, r in enumerate(results, 1):
        mins = round(r["time_spent_seconds"] / 60, 1)
        passed = "O'tdi" if r["score_percentage"] >= 60 else "O'tmadi"
        row_data = [
            idx,
            r["full_name"] or "Foydalanuvchi",
            r["telegram_id"],
            r["correct_answers"],
            r["wrong_answers"],
            r["total_questions"],
            f"{r['score_percentage']}%",
            passed,
            mins,
            r["completed_at"]
        ]
        ws1.append(row_data)
        current_row = ws1.max_row
        for col_idx in range(1, len(row_data) + 1):
            c = ws1.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 3, 4, 5, 6, 7, 8, 9, 10]:
                c.alignment = center_align

    # --- SHEET 2: REYTING (LEADERBOARD) ---
    ws2 = wb.create_sheet(title="Foydalanuvchilar Reytingi")
    headers2 = ["O'rin", "Ism Familiya", "Telegram ID", "Username", "Urinishlar soni", "Eng yuqori ball (%)", "O'rtacha ball (%)"]
    ws2.append(headers2)

    for col_idx, _ in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    analytics = get_detailed_analytics()
    for idx, u in enumerate(analytics["leaderboard"], 1):
        row_data = [
            idx,
            u["full_name"] or "—",
            u["telegram_id"],
            f"@{u['username']}" if u.get("username") else "—",
            u["attempts_count"],
            f"{u['best_score']}%",
            f"{u['avg_score']}%"
        ]
        ws2.append(row_data)
        current_row = ws2.max_row
        for col_idx in range(1, len(row_data) + 1):
            c = ws2.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx != 2:
                c.alignment = center_align

    # --- SHEET 3: QIYIN SAVOLLAR TAHLILI ---
    ws3 = wb.create_sheet(title="Qiyin Savollar Tahlili")
    headers3 = ["№", "Savol matni", "Bo'lim", "Javob berilgan", "Xato javoblar", "Xato foizi (%)"]
    ws3.append(headers3)

    for col_idx, _ in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for idx, q in enumerate(analytics["hardest_questions"], 1):
        row_data = [
            idx,
            q["question"],
            q["category"],
            q["total_answered"],
            q["wrong_count"],
            f"{q['error_rate']}%"
        ]
        ws3.append(row_data)
        current_row = ws3.max_row
        for col_idx in range(1, len(row_data) + 1):
            c = ws3.cell(row=current_row, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 3, 4, 5, 6]:
                c.alignment = center_align

    # Auto-adjust column widths for all sheets
    for sheet in [ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()

