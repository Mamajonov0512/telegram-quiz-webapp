from io import BytesIO
from fastapi.testclient import TestClient
import openpyxl
from api import app
from database import init_db, save_test_result, add_allowed_user

def test_analytics_and_excel():
    print("[*] Natijalar tahlili va Excel hisoboti testlari boshlanmoqda...")
    init_db()

    client = TestClient(app)
    admin_key = "admin12345"

    # Add dummy test results to have rich data
    dummy_users = [
        (1001, "Jasur Aliyev", 50, 48, 2, 96.0, 1500),
        (1002, "Malika Karimova", 50, 42, 8, 84.0, 1800),
        (1003, "Bobur Mirzayev", 50, 35, 15, 70.0, 2100),
        (1004, "Sardor Rahimov", 50, 22, 28, 44.0, 2400),
        (1005, "Nodira Umarova", 50, 18, 32, 36.0, 2600),
    ]

    sample_answers_json = """[
        {"id": 1, "question": "Savol 1", "category": "Tarix", "is_correct": true},
        {"id": 2, "question": "Savol 2", "category": "Tarix", "is_correct": false},
        {"id": 3, "question": "Savol 3", "category": "Fizika", "is_correct": false},
        {"id": 4, "question": "Savol 4", "category": "Matematika", "is_correct": true}
    ]"""

    for uid, name, tot, cor, wrg, pct, t_sec in dummy_users:
        add_allowed_user(uid, name, f"user_{uid}")
        save_test_result(
            telegram_id=uid,
            full_name=name,
            username=f"user_{uid}",
            total_questions=tot,
            correct_answers=cor,
            wrong_answers=wrg,
            score_percentage=pct,
            time_spent_seconds=t_sec,
            answers_json=sample_answers_json
        )

    # 1. Test Admin Analytics endpoint
    res = client.get("/api/admin/analytics", headers={"X-Admin-Key": admin_key})
    assert res.status_code == 200, f"Analytics API failed: {res.status_code}"
    data = res.json()["analytics"]

    assert data["total_tests"] >= 5
    assert data["unique_users"] >= 5
    assert "score_distribution" in data
    assert "leaderboard" in data
    assert len(data["leaderboard"]) >= 5
    assert "hardest_questions" in data
    assert "category_performance" in data
    print(f"[+] 1. Admin analitika API muvaffaqiyatli ishladi (Jami testlar: {data['total_tests']}, O'rtacha ball: {data['avg_score']}%)")
    print(f"    - Pass rate: {data['pass_rate']}%")
    print(f"    - Leaderboard top: {data['leaderboard'][0]['full_name']} ({data['leaderboard'][0]['best_score']}%)")

    # 2. Test User Personal Analytics endpoint
    user_res = client.get(f"/api/user/analytics?telegram_id=1001")
    assert user_res.status_code == 200
    u_data = user_res.json()["analytics"]
    assert u_data["total_attempts"] >= 1
    assert u_data["best_score"] == 96.0
    assert len(u_data["category_performance"]) > 0
    print(f"[+] 2. Foydalanuvchi shaxsiy tahlili muvaffaqiyatli ishladi (Best: {u_data['best_score']}%, Kategoriyalar: {len(u_data['category_performance'])})")

    # 3. Test Excel report export
    excel_res = client.get(f"/api/admin/export/excel?key={admin_key}")
    assert excel_res.status_code == 200
    assert excel_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    # Verify openpyxl can read the Excel workbook
    wb = openpyxl.load_workbook(BytesIO(excel_res.content))
    expected_sheets = ["Barcha Natijalar", "Foydalanuvchilar Reytingi", "Qiyin Savollar Tahlili"]
    for s_name in expected_sheets:
        assert s_name in wb.sheetnames, f"Sheet '{s_name}' not found in Excel!"
    
    ws1 = wb["Barcha Natijalar"]
    assert ws1.max_row > 1, "Results sheet is empty"
    ws2 = wb["Foydalanuvchilar Reytingi"]
    assert ws2.max_row > 1, "Leaderboard sheet is empty"
    print(f"[+] 3. Excel hisoboti (.xlsx) muvaffaqiyatli yaratildi (3 ta varaq: {wb.sheetnames})")

    print("\n" + "="*50)
    print("NATIJALAR TAHLILI TESTLARI 100% MUVAFFAQIYATLI O'TDI! [OK]")
    print("="*50)

if __name__ == "__main__":
    test_analytics_and_excel()
